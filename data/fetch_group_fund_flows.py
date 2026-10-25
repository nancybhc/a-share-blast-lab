from __future__ import annotations

import csv
import json
import math
import random
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from openpyxl import load_workbook


WORKBOOK = Path(r"C:\Users\jiach\多空趋势分析结果\多空阶段_2026-10-23.xlsx")
WEBSITE_DATA = Path(
    r"C:\Users\jiach\Documents\Codex\2026-10-24\d\outputs\a-share-blast-lab\data"
)
DELIVERABLE_DIR = Path(r"C:\Users\jiach\Documents\Codex\2026-10-24\d\outputs")
TARGET_DATE = "2026-10-23"
API = (
    "https://vip.stock.finance.sina.com.cn/quotes_service/api/json_v2.php/"
    "MoneyFlow.ssl_qsfx_zjlrqs"
)
HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
    "Referer": "https://finance.sina.com.cn/",
}


@dataclass(frozen=True)
class Stock:
    code: str
    name: str
    group: str

    @property
    def symbol(self) -> str:
        if self.code.startswith("920"):
            return f"bj{self.code}"
        if self.code.startswith(("5", "6", "9")):
            return f"sh{self.code}"
        return f"sz{self.code}"


def load_groups() -> tuple[dict[str, list[Stock]], dict[str, int]]:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=True)
    groups: dict[str, list[Stock]] = {}
    expected: dict[str, int] = {}

    overview = workbook["总览"]
    for row in overview.iter_rows(min_row=3, values_only=True):
        group, count = row[0], row[1]
        if group and isinstance(count, (int, float)):
            expected[str(group).strip()] = int(count)

    for sheet in workbook.worksheets[3:]:
        members: list[Stock] = []
        for row in sheet.iter_rows(min_row=5, values_only=True):
            code = str(row[3] or "").strip()
            name = str(row[4] or "").strip()
            if code.isdigit() and name:
                members.append(Stock(code=code, name=name, group=sheet.title))
        groups[sheet.title] = members
        if expected.get(sheet.title, len(members)) != len(members):
            print(
                f"[WARN] {sheet.title}: workbook summary={expected.get(sheet.title)} "
                f"parsed={len(members)}",
                flush=True,
            )

    return groups, expected


def fetch_stock(stock: Stock) -> tuple[Stock, list[dict[str, Any]], str]:
    params = {
        "page": 1,
        "num": 60,
        "sort": "opendate",
        "asc": 0,
        "daima": stock.symbol,
    }
    last_error = ""
    for attempt in range(5):
        try:
            response = requests.get(
                API,
                params=params,
                headers=HEADERS,
                timeout=18,
            )
            response.raise_for_status()
            rows = response.json()
            if not isinstance(rows, list):
                raise ValueError(f"unexpected payload: {type(rows).__name__}")
            return stock, rows, ""
        except Exception as exc:  # noqa: BLE001
            last_error = f"{type(exc).__name__}: {exc}"
            time.sleep(0.45 * (attempt + 1) + random.random() * 0.25)
    return stock, [], last_error


def numeric(row: dict[str, Any], field: str) -> float | None:
    value = row.get(field)
    if value in (None, "", "--"):
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return number if math.isfinite(number) else None


def main() -> None:
    groups, expected = load_groups()
    unique = {
        stock.code: stock
        for members in groups.values()
        for stock in members
    }
    print(
        f"[INFO] groups={len(groups)} memberships={sum(map(len, groups.values()))} "
        f"unique={len(unique)}",
        flush=True,
    )

    raw_by_code: dict[str, list[dict[str, Any]]] = {}
    errors: dict[str, str] = {}
    completed = 0
    with ThreadPoolExecutor(max_workers=20) as executor:
        futures = [executor.submit(fetch_stock, stock) for stock in unique.values()]
        for future in as_completed(futures):
            stock, rows, error = future.result()
            if rows:
                raw_by_code[stock.code] = rows
            else:
                errors[stock.code] = error
            completed += 1
            if completed % 75 == 0 or completed == len(futures):
                print(
                    f"[INFO] fetched {completed}/{len(futures)} "
                    f"ok={len(raw_by_code)} failed={len(errors)}",
                    flush=True,
                )

    all_dates = sorted(
        {
            str(row.get("opendate"))
            for rows in raw_by_code.values()
            for row in rows
            if row.get("opendate")
        }
    )
    last_six = all_dates[-6:]
    last_five = all_dates[-5:]
    if TARGET_DATE not in all_dates:
        raise RuntimeError(f"target date {TARGET_DATE} missing from source")

    website_dir = WEBSITE_DATA
    website_dir.mkdir(parents=True, exist_ok=True)
    DELIVERABLE_DIR.mkdir(parents=True, exist_ok=True)

    stock_rows: list[dict[str, Any]] = []
    group_payload: dict[str, dict[str, Any]] = {}
    summary_rows: list[dict[str, Any]] = []

    for group, members in groups.items():
        available: list[Stock] = []
        missing: list[Stock] = []
        daily_main: dict[str, float] = {day: 0.0 for day in last_six}
        daily_total: dict[str, float] = {day: 0.0 for day in last_six}
        covered_by_day: dict[str, int] = {day: 0 for day in last_six}
        target_stocks: list[dict[str, Any]] = []

        for stock in members:
            rows = raw_by_code.get(stock.code, [])
            by_date = {str(row.get("opendate")): row for row in rows}
            if TARGET_DATE not in by_date:
                missing.append(stock)
                continue
            available.append(stock)

            for day in last_six:
                row = by_date.get(day)
                if not row:
                    continue
                main = numeric(row, "r0_net")
                total = numeric(row, "netamount")
                if main is not None:
                    daily_main[day] += main
                    covered_by_day[day] += 1
                if total is not None:
                    daily_total[day] += total

            target = by_date[TARGET_DATE]
            main = numeric(target, "r0_net")
            total = numeric(target, "netamount")
            stock_rows.append(
                {
                    "group": group,
                    "code": stock.code,
                    "name": stock.name,
                    "date": TARGET_DATE,
                    "main_net_yuan": main,
                    "main_net_yi": round(main / 1e8, 4) if main is not None else None,
                    "main_net_ratio": numeric(target, "r0_ratio"),
                    "total_net_yuan": total,
                    "total_net_yi": round(total / 1e8, 4) if total is not None else None,
                    "price": numeric(target, "trade"),
                    "change_pct": (
                        numeric(target, "changeratio") * 100
                        if numeric(target, "changeratio") is not None
                        else None
                    ),
                }
            )
            if main is not None:
                target_stocks.append(
                    {
                        "code": stock.code,
                        "name": stock.name,
                        "main_net_yi": round(main / 1e8, 4),
                    }
                )

        daily_series = [
            {
                "date": day,
                "main_net_yi": round(daily_main[day] / 1e8, 2),
                "total_net_yi": round(daily_total[day] / 1e8, 2),
                "covered": covered_by_day[day],
            }
            for day in last_six
        ]
        target_main = daily_main[TARGET_DATE]
        previous_day = last_six[-2] if len(last_six) > 1 else None
        previous_main = daily_main.get(previous_day, 0.0) if previous_day else 0.0
        cumulative_main = sum(daily_main[day] for day in last_five)
        top_positive = sorted(target_stocks, key=lambda item: item["main_net_yi"], reverse=True)[:5]
        top_negative = sorted(target_stocks, key=lambda item: item["main_net_yi"])[:5]

        group_payload[group] = {
            "group": group,
            "constituents": len(members),
            "covered": len(available),
            "missing": [s.code for s in missing],
            "coverage_pct": round(len(available) / len(members) * 100, 2) if members else 0,
            "date": TARGET_DATE,
            "main_net_yi": round(target_main / 1e8, 2),
            "main_net_change_yi": round((target_main - previous_main) / 1e8, 2),
            "five_day_main_net_yi": round(cumulative_main / 1e8, 2),
            "positive_stock_count": sum(
                1
                for item in target_stocks
                if item["main_net_yi"] > 0
            ),
            "negative_stock_count": sum(
                1
                for item in target_stocks
                if item["main_net_yi"] < 0
            ),
            "top_positive": top_positive,
            "top_negative": top_negative,
            "daily": daily_series,
        }
        summary_rows.append(
            {
                "分组": group,
                "交易日": TARGET_DATE,
                "数据源": "新浪财经 MoneyFlow",
                "口径": "主力净流入 r0_net",
                "分组成分数": len(members),
                "成功覆盖数": len(available),
                "覆盖率%": group_payload[group]["coverage_pct"],
                "当日主力净流入亿": group_payload[group]["main_net_yi"],
                "较前日变化亿": group_payload[group]["main_net_change_yi"],
                "近5日累计主力净流入亿": group_payload[group]["five_day_main_net_yi"],
                "主力净流入个股数": group_payload[group]["positive_stock_count"],
                "主力净流出个股数": group_payload[group]["negative_stock_count"],
                "缺失代码": ",".join(s.code for s in missing),
            }
        )

    retrieved_at = datetime.now().astimezone().isoformat(timespec="seconds")
    payload = {
        "source": "新浪财经 MoneyFlow.ssl_qsfx_zjlrqs",
        "definition": "r0_net，即新浪财经页面标注的“主力净流入”",
        "retrieved_at": retrieved_at,
        "date": TARGET_DATE,
        "dates": last_six,
        "groups": group_payload,
        "errors": errors,
        "expected_group_sizes": expected,
    }

    json_path = website_dir / f"group_fund_flows_{TARGET_DATE}.json"
    json_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    js_path = website_dir / f"group_fund_flows_{TARGET_DATE}.js"
    js_path.write_text(
        "window.GROUP_FUND_DATA = "
        + json.dumps(payload, ensure_ascii=False, indent=2)
        + ";\n",
        encoding="utf-8",
    )

    stock_csv = website_dir / f"stock_fund_flows_{TARGET_DATE}.csv"
    with stock_csv.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(stock_rows[0].keys()))
        writer.writeheader()
        writer.writerows(stock_rows)

    summary_csv = DELIVERABLE_DIR / f"多空分组资金流_{TARGET_DATE}.csv"
    with summary_csv.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(summary_rows[0].keys()))
        writer.writeheader()
        writer.writerows(
            sorted(summary_rows, key=lambda row: float(row["当日主力净流入亿"]), reverse=True)
        )

    print(f"[DONE] json={json_path}", flush=True)
    print(f"[DONE] js={js_path}", flush=True)
    print(f"[DONE] stock_csv={stock_csv}", flush=True)
    print(f"[DONE] summary_csv={summary_csv}", flush=True)
    print(
        "[DONE] "
        f"groups_ok={sum(1 for item in group_payload.values() if not item['missing'])}/"
        f"{len(group_payload)} stock_errors={len(errors)}",
        flush=True,
    )


if __name__ == "__main__":
    main()
