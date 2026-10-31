from __future__ import annotations

import argparse
import os
import re
import shutil
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook


ANALYSIS_DIR = Path.home() / "多空趋势分析结果"
DESKTOP_DIR = Path.home() / "Desktop" / "今日盘中" / "多空趋势"
BACKUP_ROOT = (
    Path.home()
    / "Documents"
    / "Codex"
    / "2026-10-24"
    / "d"
    / "outputs"
    / "market_group_backups"
)
EXCLUDED_GROUPS = ("黄金白银",)


def workbook_paths(target_date: str) -> list[Path]:
    names = (
        f"多空阶段_{target_date}.xlsx",
        f"多空阶段迁移日报_{target_date}.xlsx",
    )
    paths = []
    for directory in (ANALYSIS_DIR, DESKTOP_DIR):
        for name in names:
            path = directory / name
            if path.is_file():
                paths.append(path)
    return paths


def shifted_ranges(ws, deleted_rows: set[int]):
    ranges = []
    for merged in list(ws.merged_cells.ranges):
        ranges.append(
            (
                merged.min_col,
                merged.min_row,
                merged.max_col,
                merged.max_row,
            )
        )
    for merged in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(merged))
    return ranges


def restore_ranges(ws, ranges, deleted_rows: set[int]) -> None:
    ordered = sorted(deleted_rows)
    for min_col, min_row, max_col, max_row in ranges:
        before_start = sum(row < min_row for row in ordered)
        through_end = sum(row <= max_row for row in ordered)
        new_min = min_row - before_start
        new_max = max_row - through_end
        if new_min <= new_max:
            ws.merge_cells(
                start_row=new_min,
                start_column=min_col,
                end_row=new_max,
                end_column=max_col,
            )


def remove_group_rows(ws) -> int:
    rows = []
    for row_no, row in enumerate(
        ws.iter_rows(values_only=True),
        1,
    ):
        if any(value in EXCLUDED_GROUPS for value in row):
            rows.append(row_no)
    if not rows:
        return 0

    ranges = shifted_ranges(ws, set(rows))
    for row_no in sorted(rows, reverse=True):
        ws.delete_rows(row_no, 1)
    restore_ranges(ws, ranges, set(rows))
    return len(rows)


def update_group_counts(ws, group_count: int) -> None:
    if ws.title == "总览":
        cell = ws["A1"]
        if isinstance(cell.value, str):
            cell.value = re.sub(r"\d+\s*个分组", f"{group_count} 个分组", cell.value)
    if ws.title == "分组异动":
        cell = ws["A1"]
        if isinstance(cell.value, str):
            cell.value = re.sub(r"\d+\s*个分组", f"{group_count} 个分组", cell.value)
    if ws.title == "今日重点":
        for cell in ws[2]:
            if isinstance(cell.value, str):
                cell.value = re.sub(r"\d+\s*个分组", f"{group_count} 个分组", cell.value)
        if isinstance(ws["A5"].value, (int, float)):
            ws["A5"] = group_count
        if isinstance(ws["A6"].value, str):
            ws["A6"].value = re.sub(
                r"\d+\s*组分口径",
                f"{group_count} 组分口径",
                ws["A6"].value,
            )
        priority_rows = sum(
            1
            for row_no in range(10, ws.max_row + 1)
            if str(ws.cell(row_no, 1).value or "").strip()
            in {"高优先", "中优先", "低优先"}
        )
        if isinstance(ws["A8"].value, str) and "今天先看" in ws["A8"].value:
            ws["A8"] = f"今天先看这 {priority_rows} 件事"


def filter_workbook(path: Path, backup_dir: Path) -> dict[str, object]:
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup_path = backup_dir / path.name
    if not backup_path.exists():
        shutil.copy2(path, backup_path)

    workbook = load_workbook(path)
    removed_sheets = []
    removed_rows = {}
    for sheet_name in EXCLUDED_GROUPS:
        if sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            workbook.remove(sheet)
            removed_sheets.append(sheet_name)

    if "分组异动" in workbook.sheetnames:
        migration_sheet = workbook["分组异动"]
        group_count = sum(
            1
            for values in migration_sheet.iter_rows(min_row=3, values_only=True)
            if len(values) > 1 and values[0] and values[1]
        )
    else:
        summary_sheets = {"总览", "量化5破10推演", "多空趋势"}
        group_count = sum(
            sheet.title not in summary_sheets for sheet in workbook.worksheets
        )

    for sheet in workbook.worksheets:
        count = remove_group_rows(sheet)
        if count:
            removed_rows[sheet.title] = count

    if "分组异动" in workbook.sheetnames:
        migration_sheet = workbook["分组异动"]
        group_count = sum(
            1
            for values in migration_sheet.iter_rows(min_row=3, values_only=True)
            if len(values) > 1 and values[0] and values[1]
        )

    for sheet in workbook.worksheets:
        update_group_counts(sheet, group_count)

    temp_path = path.with_suffix(path.suffix + ".tmp")
    workbook.save(temp_path)
    workbook.close()
    os.replace(temp_path, path)
    return {
        "path": str(path),
        "backup": str(backup_path),
        "removed_sheets": removed_sheets,
        "removed_rows": removed_rows,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="从多空阶段 Excel 剔除指定分组")
    parser.add_argument("--date", default=datetime.now().strftime("%Y-%m-%d"))
    parser.add_argument("--backup-root", type=Path, default=BACKUP_ROOT)
    args = parser.parse_args()

    paths = workbook_paths(args.date)
    if not paths:
        raise FileNotFoundError(f"{args.date} 未找到多空阶段 Excel")
    backup_dir = args.backup_root / args.date
    for path in paths:
        result = filter_workbook(path, backup_dir)
        print(
            f"[OK] {path} sheets={result['removed_sheets']} "
            f"rows={result['removed_rows']}"
        )


if __name__ == "__main__":
    main()
