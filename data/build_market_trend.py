from __future__ import annotations

import argparse
import json
import re
import shutil
from datetime import date, datetime
from pathlib import Path

import openpyxl


DEFAULT_SOURCE_DIR = Path(r"C:\Users\jiach\Desktop\今日盘中\多空趋势")
DEFAULT_OUTPUT_DIR = Path(__file__).resolve().parent

GROUP_META = {
    "PCB": ("pcb", "PCB"),
    "CPO": ("cpo", "CPO"),
    "光上游": ("optics", "OPTICS"),
    "半导体芯片": ("semiconductor", "SEMICONDUCTOR"),
    "半导体设备": ("equipment", "SEMI CAP"),
    "半导体上游": ("semiconductor-upstream", "SEMI MATERIALS"),
    "MLCC": ("mlcc", "PASSIVE"),
    "算力": ("compute", "COMPUTE"),
    "存储": ("storage", "MEMORY"),
    "互金证券": ("fintech", "FINTECH"),
    "AI应用": ("ai-app", "AI APPS"),
    "机器人": ("robot", "ROBOTICS"),
    "药": ("pharma", "PHARMA"),
    "黄金白银": ("gold", "PRECIOUS METALS"),
    "电力": ("power", "POWER"),
    "电网设备": ("grid-equipment", "GRID EQUIPMENT"),
    "商业航天": ("aerospace", "AEROSPACE"),
}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="把最新多空阶段报告导出为网页数据。")
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_SOURCE_DIR)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--date", help="交易日 YYYY-MM-DD；默认选择最新报告。")
    return parser.parse_args()


def latest_file(directory: Path, pattern: str, target_date: str | None) -> Path:
    if target_date:
        path = directory / pattern.format(date=target_date)
        if not path.is_file():
            raise FileNotFoundError(path)
        return path
    candidates = sorted(directory.glob(pattern.format(date="*")), key=lambda path: path.stat().st_mtime)
    if not candidates:
        raise FileNotFoundError(f"未找到 {directory / pattern.format(date='*')}")
    return candidates[-1]


def date_text(value: object) -> str:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value or "")[:10]


def read_group_rows(workbook: openpyxl.Workbook) -> list[dict[str, object]]:
    sheet = workbook["分组异动"]
    rows: list[dict[str, object]] = []
    for values in sheet.iter_rows(min_row=3, values_only=True):
        if not values[0] or not values[1]:
            continue
        rows.append(
            {
                "name": str(values[0]),
                "date": date_text(values[1]),
                "total": int(values[2]),
                "bull": int(values[3]),
                "turnBull": int(values[4]),
                "turnBear": int(values[5]),
                "bear": int(values[6]),
                "bullPct": float(values[7]),
                "newBull": int(values[12]),
                "confirm": int(values[13]),
                "direct": int(values[14]),
                "positive": int(values[17]),
                "negative": int(values[18]),
                "state": str(values[19] or ""),
                "change": str(values[20] or ""),
                "action": str(values[21] or ""),
            }
        )
    return rows


def read_histories(workbook: openpyxl.Workbook, names: set[str]) -> dict[str, list[dict[str, object]]]:
    sheet = workbook["多空趋势"]
    records: dict[str, list[dict[str, object]]] = {}
    for values in sheet.iter_rows(values_only=True):
        name = str(values[0] or "").strip()
        if name not in names or not values[1]:
            continue
        records.setdefault(name, []).append(
            {
                "date": date_text(values[1]),
                "bull": int(values[2]),
                "turnBull": int(values[3]),
                "turnBear": int(values[4]),
                "bear": int(values[5]),
                "bullPct": float(values[7]),
            }
        )
    for rows in records.values():
        rows.sort(key=lambda item: item["date"])
    return records


def phase_for(row: dict[str, object], score: float) -> tuple[str, str, str]:
    if row["direct"] > 0:
        return "空→多跳级", "修复加速", "lime"
    if score >= 15:
        return "转多兑现", "扩散确认", "lime"
    if score > 0:
        return "温和修复", "边际改善", "amber"
    if score <= -15:
        return "转空加速", "风险升高", "red"
    if score < 0:
        return "结构承压", "等待止跌", "cyan"
    return "结构分化", "临界观察", "amber"


def make_thesis(row: dict[str, object]) -> str:
    return (
        f"多头 {row['bull']}/{row['total']}，空头 {row['bear']}；"
        f"空转多候选 {row['turnBull']}，今日转多确认 {row['confirm']}，反向转空 {row['turnBear']}。"
        f"{row['change']}，{row['action']}。"
    )


def build_payload(cycle_path: Path, migration_path: Path) -> dict[str, object]:
    cycle = openpyxl.load_workbook(cycle_path, read_only=True, data_only=True)
    migration = openpyxl.load_workbook(migration_path, read_only=True, data_only=True)
    rows = read_group_rows(migration)
    histories = read_histories(cycle, {row["name"] for row in rows})

    groups: list[dict[str, object]] = []
    for row in rows:
        name = str(row["name"])
        group_id, code = GROUP_META.get(name, (re.sub(r"\W+", "-", name).lower(), name))
        raw_score = (
            row["newBull"] + row["confirm"] * 2 + row["direct"] * 3 - row["negative"]
        ) / row["total"] * 100
        score = round(max(-50.0, min(80.0, raw_score)), 1)
        phase, badge, tone = phase_for(row, score)
        groups.append(
            {
                **row,
                "id": group_id,
                "code": code,
                "score": score,
                "fulfilled": row["confirm"],
                "netBias": row["bull"] - row["bear"],
                "phase": phase,
                "badge": badge,
                "tone": tone,
                "setup": str(row["turnBull"]),
                "setupCaption": "空转多候选",
                "confirmValue": str(row["confirm"]),
                "confirmCaption": "今日转多确认",
                "thesis": make_thesis(row),
                "history": histories.get(name, [])[-7:],
            }
        )

    groups.sort(key=lambda group: (-group["score"], -group["bullPct"], str(group["name"])))
    visible_dates = sorted({item["date"] for group in groups for item in group["history"]})[-5:]
    series = []
    for name, group_id, color in (
        ("机器人", "robot", "#a95b0f"),
        ("半导体芯片", "semiconductor", "#087a88"),
    ):
        history = next(group["history"] for group in groups if group["name"] == name)
        by_date = {item["date"]: item for item in history}
        series.append(
            {
                "id": group_id,
                "name": name,
                "color": color,
                "values": [int(by_date[day]["turnBull"]) if day in by_date else 0 for day in visible_dates],
            }
        )

    title = cycle["总览"].cell(1, 1).value or ""
    runtime_match = re.search(r"运行时间\s+([0-9-]+\s+[0-9:]+)", str(title))
    summary_lines = [str(cycle["多空趋势"].cell(2, 1).value or "")]
    summary_date = groups[0]["date"]
    detailed_codes: set[str] = set()
    for sheet in cycle.worksheets[3:]:
        for values in sheet.iter_rows(min_row=4, values_only=True):
            code = str(values[3] or "").strip() if len(values) > 3 else ""
            if re.fullmatch(r"\d{6}", code):
                detailed_codes.add(code)

    positive_groups = sum(1 for group in groups if group["score"] > 0)
    negative_groups = sum(1 for group in groups if group["score"] < 0)
    summary_text = "\n".join(
        str(values[0])
        for values in cycle["总览"].iter_rows(values_only=True)
        if values and values[0] is not None
    )
    generated_at = ""
    if runtime_match:
        generated_at = runtime_match.group(1)
    else:
        generated_at = datetime.fromtimestamp(cycle_path.stat().st_mtime).strftime("%Y-%m-%d %H:%M")

    payload = {
        "dataDate": summary_date,
        "generatedAt": generated_at,
        "isIntraday": any(word in summary_text for word in ("盘中", "午间")),
        "sourceReport": cycle_path.name,
        "sourcePath": str(cycle_path),
        "sourceBrief": f"assets/多空阶段简报_{summary_date}.pptx",
        "summary": "\n".join(line for line in summary_lines if line),
        "overview": {
            "groups": len(groups),
            "uniqueStocks": len(detailed_codes),
            "positiveGroups": positive_groups,
            "negativeGroups": negative_groups,
            "positiveTransitions": sum(int(group["positive"]) for group in groups),
            "negativeTransitions": sum(int(group["negative"]) for group in groups),
        },
        "groups": groups,
        "transition": {"labels": [day[5:].replace("-", ".") for day in visible_dates], "series": series},
    }
    cycle.close()
    migration.close()
    return payload


def main() -> None:
    args = parse_args()
    cycle_path = latest_file(args.source_dir, "多空阶段_{date}.xlsx", args.date)
    report_date = args.date or re.search(r"(\d{4}-\d{2}-\d{2})", cycle_path.name).group(1)
    migration_path = latest_file(args.source_dir, "多空阶段迁移日报_{date}.xlsx", report_date)
    brief_path = latest_file(args.source_dir, "多空阶段简报_{date}.pptx", report_date)

    payload = build_payload(cycle_path, migration_path)
    output_dir = args.output_dir.resolve()
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "market_trend_latest.json"
    js_path = output_dir / "market_trend_latest.js"
    json_text = json.dumps(payload, ensure_ascii=False, indent=2)
    json_path.write_text(json_text + "\n", encoding="utf-8")
    js_path.write_text(
        "window.MARKET_TREND_DATA = " + json_text + ";\n",
        encoding="utf-8",
    )

    assets_dir = output_dir.parent / "assets"
    assets_dir.mkdir(parents=True, exist_ok=True)
    shutil.copy2(brief_path, assets_dir / brief_path.name)
    print(f"[OK] {json_path}")
    print(f"[OK] {js_path}")
    print(
        f"[OK] {payload['dataDate']} / {payload['generatedAt']} / "
        f"{payload['overview']['groups']} groups / {payload['overview']['uniqueStocks']} stocks"
    )


if __name__ == "__main__":
    main()
